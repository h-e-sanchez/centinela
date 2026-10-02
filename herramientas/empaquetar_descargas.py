"""Arma los descargables de cada reporte de la vitrina a partir de las fuentes del repo.

Por reporte (`reportes/catalogo.json`: campos `modelo`, `datos` y `prefijo`):

- <prefijo>-datos-csv.zip  los CSV de `datos` + LEEME.txt (diccionario de columnas)
- <prefijo>-datos.xlsx     una hoja por CSV y una hoja diccionario
- <prefijo>-pbip.zip       el proyecto PBIP de `modelo`, sin la caché ni la configuración local

El diccionario vive en reportes/<slug>/diccionario.json. Los zip son reproducibles (entradas
ordenadas, fecha fija). El .pbix no se genera aquí: se exporta a mano desde Power BI Desktop
(Archivo → Guardar como).

Uso (requiere openpyxl, en requirements-dev.txt):
    python herramientas/empaquetar_descargas.py                 # todos los reportes
    python herramientas/empaquetar_descargas.py --reporte workforce
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import zipfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
REPORTES = RAIZ / "reportes"
FECHA_FIJA = (2026, 1, 1, 0, 0, 0)
EXCLUIR_PBIP = {"cache.abf", "localSettings.json"}
ENTERO = re.compile(r"-?\d+")
DECIMAL = re.compile(r"-?\d+\.\d+")


def catalogo() -> list[dict]:
    return json.loads((REPORTES / "catalogo.json").read_text(encoding="utf-8"))["reportes"]


def reporte(slug: str) -> dict:
    for rep in catalogo():
        if rep["slug"] == slug:
            return rep
    raise SystemExit(f"no existe el reporte {slug!r} en reportes/catalogo.json")


def destino(rep: dict) -> Path:
    return REPORTES / rep["slug"] / "descargas"


def carpeta_datos(rep: dict) -> Path:
    return RAIZ / rep["datos"]


def carpeta_modelo(rep: dict) -> Path:
    return RAIZ / rep["modelo"]


def csvs(rep: dict) -> list[Path]:
    return sorted(carpeta_datos(rep).glob("*.csv"))


def diccionario(rep: dict) -> dict:
    return json.loads((REPORTES / rep["slug"] / "diccionario.json").read_text(encoding="utf-8"))


def _escribir_zip(ruta: Path, entradas: list[tuple[str, bytes]]) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as zf:
        for nombre, datos in sorted(entradas):
            info = zipfile.ZipInfo(nombre, date_time=FECHA_FIJA)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zf.writestr(info, datos)


def leeme(rep: dict) -> str:
    dic = diccionario(rep)
    ancho = max(len(c) for c in dic["columnas"]) + 2
    columnas = "\n".join(f"  {c:<{ancho}}{d}" for c, d in dic["columnas"].items())
    return dic["leeme"].format(columnas=columnas)


def entradas_csv(rep: dict) -> list[tuple[str, bytes]]:
    entradas = [("LEEME.txt", leeme(rep).encode("utf-8"))]
    for ruta in csvs(rep):
        entradas.append((ruta.name, ruta.read_bytes()))
    return entradas


def entradas_pbip(rep: dict) -> list[tuple[str, bytes]]:
    base = carpeta_modelo(rep)
    entradas = []
    for ruta in sorted(base.rglob("*")):
        if ruta.is_file() and ruta.name not in EXCLUIR_PBIP:
            entradas.append((ruta.relative_to(base).as_posix(), ruta.read_bytes()))
    return entradas


def leer_csv(ruta: Path) -> list[list[str]]:
    with ruta.open(encoding="utf-8", newline="") as f:
        return list(csv.reader(f))


def tipar(valor: str) -> int | float | str:
    """Número si el texto es numérico; así Excel no guarda cifras como texto."""
    if ENTERO.fullmatch(valor):
        return int(valor)
    if DECIMAL.fullmatch(valor):
        return float(valor)
    return valor


def escribir_xlsx(rep: dict, ruta: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    libro = Workbook()
    libro.remove(libro.active)
    for csv_ruta in csvs(rep):
        hoja = libro.create_sheet(csv_ruta.stem[:31])
        filas = leer_csv(csv_ruta)
        hoja.append(filas[0])
        for fila in filas[1:]:
            hoja.append([tipar(v) for v in fila])
        for celda in hoja[1]:
            celda.font = Font(bold=True)
        hoja.freeze_panes = "A2"
    dic = libro.create_sheet("diccionario")
    dic.append(["columna", "descripcion"])
    for columna, descripcion in diccionario(rep)["columnas"].items():
        dic.append([columna, descripcion])
    libro.properties.creator = "centinela"
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)


def empaquetar(rep: dict) -> None:
    salida = destino(rep)
    prefijo = rep["prefijo"]
    _escribir_zip(salida / f"{prefijo}-datos-csv.zip", entradas_csv(rep))
    _escribir_zip(salida / f"{prefijo}-pbip.zip", entradas_pbip(rep))
    escribir_xlsx(rep, salida / f"{prefijo}-datos.xlsx")
    for ruta in sorted(salida.iterdir()):
        print(f"{rep['slug']:20} {ruta.name:30} {ruta.stat().st_size / 1024:8.1f} KB")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--reporte", help="slug del reporte (por defecto, todos)")
    args = parser.parse_args()
    for rep in [reporte(args.reporte)] if args.reporte else catalogo():
        empaquetar(rep)


if __name__ == "__main__":
    main()
